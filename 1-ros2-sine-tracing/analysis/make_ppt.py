#!/usr/bin/env python3
"""
Build the presentation (PowerPoint) for the guide, from the shared SECTIONS.
Each content slide carries speaker notes (what to say + likely questions).

Run:  python3 analysis/make_ppt.py   ->  docs/R_sine_presentation.pptx
"""
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PKG = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(PKG, "analysis"))
import deliverable_data as D   # noqa: E402

FIG = os.path.join(PKG, "figures")
OUT = os.path.join(PKG, "docs", "R_sine_presentation.pptx")

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0xD6, 0x27, 0x28)
GREY = RGBColor(0x44, 0x44, 0x44)
EMU_W, EMU_H = Inches(13.333), Inches(7.5)


def _title_bar(slide, text):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.9))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = NAVY
    return box


def _bullets(slide, bullets, left, top, width, height, size=18):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame; tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + b
        p.font.size = Pt(size); p.font.color.rgb = GREY
        p.space_after = Pt(8)
    return box


def _notes(slide, sec):
    notes = slide.notes_slide.notes_text_frame
    notes.text = "WHAT TO SAY:\n" + sec.get("say", "")
    for q, a in sec.get("questions", []):
        p = notes.add_paragraph(); p.text = f"\nQ: {q}\nA: {a}"


def _add_table(slide, rows, left, top):
    header = ["i", "alpha (deg)", "a (m)", "d (m)", "theta off (deg)"]
    n = len(rows) + 1
    tbl = slide.shapes.add_table(n, 5, left, top, Inches(7.5), Inches(0.4 * n)).table
    for c, h in enumerate(header):
        cell = tbl.cell(0, c); cell.text = h
        cell.text_frame.paragraphs[0].font.size = Pt(14)
        cell.text_frame.paragraphs[0].font.bold = True
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c); cell.text = str(val)
            cell.text_frame.paragraphs[0].font.size = Pt(13)


def build():
    prs = Presentation()
    prs.slide_width = EMU_W; prs.slide_height = EMU_H
    blank = prs.slide_layouts[6]

    for i, sec in enumerate(D.SECTIONS):
        slide = prs.slides.add_slide(blank)
        fig = sec.get("figure")
        img = os.path.join(FIG, fig) if fig else None

        if i == 0:
            # title slide
            t = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.6))
            p = t.text_frame.paragraphs[0]; p.text = sec["title"]
            p.font.size = Pt(40); p.font.bold = True; p.font.color.rgb = NAVY
            p.alignment = PP_ALIGN.CENTER; t.text_frame.word_wrap = True
            s = slide.shapes.add_textbox(Inches(0.8), Inches(3.9), Inches(11.7), Inches(0.8))
            sp = s.text_frame.paragraphs[0]; sp.text = sec.get("subtitle", "") or ""
            sp.font.size = Pt(20); sp.font.color.rgb = ACCENT; sp.alignment = PP_ALIGN.CENTER
            if img:
                slide.shapes.add_picture(img, Inches(4.4), Inches(4.7), height=Inches(2.4))
            _notes(slide, sec)
            continue

        _title_bar(slide, sec["title"])
        if sec.get("subtitle"):
            sb = slide.shapes.add_textbox(Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.5))
            sp = sb.text_frame.paragraphs[0]; sp.text = sec["subtitle"]
            sp.font.size = Pt(16); sp.font.italic = True; sp.font.color.rgb = ACCENT

        top = Inches(1.7)
        if sec.get("table"):
            _bullets(slide, sec["bullets"], Inches(0.5), top, Inches(12.3), Inches(2.2), size=17)
            _add_table(slide, sec["table"], Inches(2.8), Inches(4.1))
        elif img:
            _bullets(slide, sec["bullets"], Inches(0.5), top, Inches(6.2), Inches(5.2))
            slide.shapes.add_picture(img, Inches(6.9), Inches(1.7), width=Inches(6.1))
        else:
            _bullets(slide, sec["bullets"], Inches(0.7), top, Inches(11.9), Inches(5.0), size=20)
        _notes(slide, sec)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print("wrote", OUT, f"({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
